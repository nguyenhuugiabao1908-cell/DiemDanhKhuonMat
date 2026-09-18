# StudentAttendance/gui/history.py
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from pathlib import Path
from datetime import datetime, timedelta

try:
    from tkcalendar import DateEntry
except ImportError:
    pass

from utils.database import DatabaseManager
from utils.attendance_manager import AttendanceManager
from utils.helpers import modern_theme_colors

class HistoryFrame:
    def __init__(self, parent: tk.Widget, db: DatabaseManager, base_dir: Path) -> None:
        self.parent = parent
        self.db = db
        self.base_dir = base_dir
        self.colors = modern_theme_colors()
        self.attendance_manager = AttendanceManager()
        self.tree = None
        
        self.search_var = tk.StringVar()
        self.start_date_var = tk.StringVar()
        self.end_date_var = tk.StringVar()
        
        self._build_ui()
        self.load_data()

    def _build_ui(self):
        wrapper = tk.Frame(self.parent, bg="white", highlightthickness=1, highlightbackground=self.colors["border"])
        wrapper.pack(fill="both", expand=True)

        # ==========================================
        # 1. HEADER & TÌM KIẾM CƠ BẢN
        # ==========================================
        header = tk.Frame(wrapper, bg="white")
        header.pack(fill="x", padx=20, pady=(20, 10))

        tk.Label(
            header, text="Lịch sử điểm danh", bg="white", 
            fg=self.colors["text"], font=("Segoe UI", 16, "bold")
        ).pack(side="left")

        search_frame = tk.Frame(header, bg="white")
        search_frame.pack(side="left", padx=(30, 0))
        tk.Label(search_frame, text="🔍 Tìm MSSV/Tên:", bg="white", font=("Segoe UI", 10, "bold"), fg=self.colors["muted"]).pack(side="left", padx=(0, 5))
        search_box = ttk.Entry(search_frame, textvariable=self.search_var, font=("Segoe UI", 10), width=25)
        search_box.pack(side="left", ipady=3)
        self.search_var.trace_add("write", lambda *_: self.load_data())

        # --- CÁC NÚT BẤM ---
        tk.Button(header, text="Xuất Excel", command=self.export_excel, bg="#107c41", fg="white", bd=0, padx=15, pady=6, cursor="hand2", font=("Segoe UI", 10, "bold")).pack(side="right", padx=(10, 0))
        tk.Button(header, text="Xuất PDF", command=self.export_pdf, bg="#b91c1c", fg="white", bd=0, padx=15, pady=6, cursor="hand2", font=("Segoe UI", 10, "bold")).pack(side="right", padx=(10, 0))
        tk.Button(header, text="Xóa lượt chọn", command=self.delete_attendance, bg="#dc2626", fg="white", bd=0, padx=15, pady=6, cursor="hand2", font=("Segoe UI", 10, "bold")).pack(side="right", padx=(10, 0))
        tk.Button(header, text="Làm mới", command=self.load_data, bg="#2563eb", fg="white", bd=0, padx=15, pady=6, cursor="hand2", font=("Segoe UI", 10, "bold")).pack(side="right")

        # ==========================================
        # 2. KHU VỰC BỘ LỌC THỜI GIAN & THỐNG KÊ 
        # ==========================================
        stats_box = tk.Frame(wrapper, bg="#f8fafc", highlightthickness=1, highlightbackground=self.colors["border"])
        stats_box.pack(fill="x", padx=20, pady=(0, 15))

        filter_row = tk.Frame(stats_box, bg="#f8fafc")
        filter_row.pack(fill="x", padx=15, pady=(15, 10))

        # --- BỘ LỌC LỚP ---
        tk.Label(filter_row, text="Lớp:", bg="#f8fafc", font=("Segoe UI", 10, "bold")).pack(side="left", padx=(0, 5))
        
        all_st = self.db.get_students()
        class_list = ["Tất cả"] + sorted(list(set(row["class_name"] for row in all_st if row["class_name"])))
        
        self.class_filter = ttk.Combobox(filter_row, values=class_list, state="readonly", width=12, font=("Segoe UI", 10))
        self.class_filter.current(0)
        self.class_filter.pack(side="left", padx=(0, 20))
        self.class_filter.bind("<<ComboboxSelected>>", lambda e: self.load_data())

        tk.Label(filter_row, text="Trạng thái:", bg="#f8fafc", font=("Segoe UI", 10, "bold")).pack(side="left", padx=(0, 5))
        self.status_filter = ttk.Combobox(filter_row, values=["Tất cả", "Có mặt", "Đi muộn", "Vắng mặt"], state="readonly", width=12, font=("Segoe UI", 10))
        self.status_filter.current(0)
        self.status_filter.pack(side="left")
        self.status_filter.bind("<<ComboboxSelected>>", lambda e: self.load_data())

        # Lấy ngày hiện tại
        today_str = datetime.now().strftime("%Y-%m-%d")

        tk.Label(filter_row, text="Từ ngày:", bg="#f8fafc", font=("Segoe UI", 10)).pack(side="left", padx=(30, 5))
        self.start_date_entry = DateEntry(
            filter_row, textvariable=self.start_date_var, font=("Segoe UI", 10), width=12,
            date_pattern='yyyy-mm-dd', background=self.colors["primary"], 
            foreground='white', borderwidth=1, cursor="hand2",
            showweeknumbers=False
        )
        self.start_date_entry.pack(side="left")
        self.start_date_var.set(today_str) # Tự động điền ngày hiện tại

        tk.Label(filter_row, text="Đến ngày:", bg="#f8fafc", font=("Segoe UI", 10)).pack(side="left", padx=(15, 5))
        self.end_date_entry = DateEntry(
            filter_row, textvariable=self.end_date_var, font=("Segoe UI", 10), width=12,
            date_pattern='yyyy-mm-dd', background=self.colors["primary"], 
            foreground='white', borderwidth=1, cursor="hand2",
            showweeknumbers=False
        )
        self.end_date_entry.pack(side="left")
        self.end_date_var.set(today_str) # Tự động điền ngày hiện tại

        tk.Button(filter_row, text="Lọc dữ liệu", command=self.load_data, bg=self.colors["primary"], fg="white", bd=0, padx=15, pady=4, cursor="hand2", font=("Segoe UI", 9, "bold")).pack(side="left", padx=(15, 0))

        # --- Bảng chỉ số thống kê ---
        metrics_row = tk.Frame(stats_box, bg="#f8fafc")
        metrics_row.pack(fill="x", padx=15, pady=(0, 15))

        self.lbl_total = tk.Label(metrics_row, text="Tổng lượt học theo TKB: 0", bg="#f8fafc", fg="#475569", font=("Segoe UI", 11, "bold"))
        self.lbl_total.pack(side="left", padx=(0, 30))
        self.lbl_present = tk.Label(metrics_row, text="Có mặt: 0", bg="#f8fafc", fg="#16a34a", font=("Segoe UI", 11, "bold"))
        self.lbl_present.pack(side="left", padx=(0, 30))
        self.lbl_absent = tk.Label(metrics_row, text="Vắng: 0", bg="#f8fafc", fg="#dc2626", font=("Segoe UI", 11, "bold"))
        self.lbl_absent.pack(side="left", padx=(0, 30))
        self.lbl_rate = tk.Label(metrics_row, text="Tỷ lệ tham dự: 0%", bg="#f8fafc", fg="#7c3aed", font=("Segoe UI", 11, "bold"))
        self.lbl_rate.pack(side="left")

        # ==========================================
        # 3. BẢNG DỮ LIỆU
        # ==========================================
        body = tk.Frame(wrapper, bg="white")
        body.pack(fill="both", expand=True, padx=20, pady=(0, 20))

        columns = ("student_id", "full_name", "class_name", "subject", "date", "time", "status")
        self.tree = ttk.Treeview(body, columns=columns, show="headings", height=20)

        # BẮT SỰ KIỆN PHÍM TẮT CTRL + A
        self.tree.bind("<Control-a>", self.select_all)
        self.tree.bind("<Control-A>", self.select_all) 

        headings = ["MSSV", "Họ tên", "Lớp", "Môn học", "Ngày", "Giờ", "Trạng thái"]
        widths = [90, 160, 80, 180, 100, 100, 120]

        for col, title, w in zip(columns, headings, widths):
            self.tree.heading(col, text=title)
            self.tree.column(col, width=w, anchor="center")
            
        scrollbar = ttk.Scrollbar(body, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar.set)
        
        self.tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        self.tree.tag_configure("absent", foreground="#dc2626")

    # ==========================================
    # CÁC HÀM XỬ LÝ LOGIC 
    # ==========================================
    def load_data(self):
        if self.tree is None: return
        for item in self.tree.get_children():
            self.tree.delete(item)
            
        keyword = self.search_var.get().strip().lower()
        start_date_str = self.start_date_var.get().strip()
        end_date_str = self.end_date_var.get().strip()
        status_filter = self.status_filter.get()
        selected_class = self.class_filter.get()
        
        raw_students = self.db.get_students()
        
        # Cập nhật danh sách Lớp động vào Combobox
        class_list = ["Tất cả"] + sorted(list(set(row["class_name"] for row in raw_students if row["class_name"])))
        if self.class_filter['values'] != tuple(class_list):
            self.class_filter['values'] = class_list
        
        if selected_class == "Tất cả":
            all_students = raw_students
        else:
            all_students = [s for s in raw_students if s["class_name"] == selected_class]
            
        student_dict = {s["student_id"]: s for s in all_students}
        all_history = self.db.attendance_history()
        
        display_rows = []
        total_present = 0
        total_expected = 0
        
        # Xử lý mốc thời gian (Mặc định lấy 7 ngày qua để tránh lag)
        if start_date_str: 
            start_dt = datetime.strptime(start_date_str, "%Y-%m-%d").date()
        else: 
            start_dt = (datetime.now() - timedelta(days=7)).date()
            
        if end_date_str: 
            end_dt = datetime.strptime(end_date_str, "%Y-%m-%d").date()
        else: 
            end_dt = datetime.now().date()

        history_set = set()
        
        # 1. TRÍCH XUẤT NHỮNG LƯỢT CÓ MẶT TỪ DATABASE
        for r in all_history:
            s_id = r["student_id"]
            if s_id not in student_dict: continue
            
            d_str = r["attendance_date"]
            d_obj = datetime.strptime(d_str, "%Y-%m-%d").date()
            
            if start_dt <= d_obj <= end_dt:
                total_present += 1
                sub_name = r["subject_name"]
                history_set.add((s_id, d_str, sub_name))
                
                status_label = "Đi muộn" if r["status"] == "Late" else "Có mặt"
                if status_filter in ["Tất cả", status_label]:
                    display_rows.append((s_id, r["full_name"], r["class_name"], sub_name, d_str, r["attendance_time"], status_label, "present"))

        # 2. SUY LUẬN NHỮNG LƯỢT VẮNG MẶT TỪ THỜI KHÓA BIỂU
        subjects_map = {sub["subject_id"]: sub["subject_name"] for sub in self.db.get_subjects()}
        
        for s_id, student in student_dict.items():
            cls = student["class_name"].strip()
            schedules = self.db.get_schedules(cls)
            
            curr = start_dt
            while curr <= end_dt:
                day_idx = curr.weekday()
                # Tìm xem ngày này lớp có môn gì không
                daily_subs = [sch for sch in schedules if sch["day_of_week"] == day_idx]
                date_str = curr.strftime("%Y-%m-%d")
                
                for sch in daily_subs:
                    total_expected += 1
                    sub_name = subjects_map.get(sch["subject_id"], sch["subject_id"])
                    
                    # Nếu lịch học yêu cầu có mặt mà DB không có dữ liệu -> Vắng
                    if (s_id, date_str, sub_name) not in history_set:
                        if status_filter in ["Tất cả", "Vắng mặt"]:
                            display_rows.append((s_id, student["full_name"], student["class_name"], sub_name, date_str, "--:--", "Absent", "absent"))
                
                curr += timedelta(days=1)
                
        # 3. SẮP XẾP VÀ HIỂN THỊ
        display_rows.sort(key=lambda x: x[4], reverse=True)
        
        for row in display_rows:
            display_tuple = row[:7]
            tag = row[7]
            if keyword and keyword not in str(display_tuple).lower():
                continue
            self.tree.insert("", "end", values=display_tuple, tags=(tag,))
            
        self.calculate_statistics(total_expected, total_present)

    def calculate_statistics(self, expected_total, actual_present):
        absent_count = max(0, expected_total - actual_present)
        rate = (actual_present / expected_total * 100) if expected_total > 0 else 0.0

        self.lbl_total.config(text=f"Tổng lượt học theo TKB: {expected_total}")
        self.lbl_present.config(text=f"Có mặt: {actual_present}")
        self.lbl_absent.config(text=f"Vắng: {absent_count}")
        self.lbl_rate.config(text=f"Tỷ lệ tham dự: {rate:.1f}%")

    def delete_attendance(self):
        selected = self.tree.selection()
        if not selected:
            messagebox.showwarning("Cảnh báo", "Vui lòng bấm chọn ít nhất một dòng để xóa.")
            return
            
        valid_records = []
        for item in selected:
            values = self.tree.item(item, "values")
            student_id, full_name, class_name, subject, date, time, status = values
            
            if status != "Absent" and time != "--:--":
                valid_records.append((student_id, date, time, full_name))

        if not valid_records:
            messagebox.showinfo("Thông báo", "Các dòng bạn chọn đều là trạng thái Vắng mặt do hệ thống tự động suy luận ra, không thể xóa!\nBạn chỉ có thể xóa các lượt Có mặt.")
            return

        if len(valid_records) == 1:
            msg = f"Bạn có chắc muốn HỦY điểm danh của sinh viên:\n{valid_records[0][3]}\nLúc: {valid_records[0][2]} - {valid_records[0][1]}?"
        else:
            msg = f"Bạn có chắc chắn muốn HỦY {len(valid_records)} lượt điểm danh hợp lệ đã chọn không?"

        if not messagebox.askyesno("Xác nhận", msg):
            return
            
        try:
            for record in valid_records:
                student_id, date, time, _ = record
                self.db.execute("DELETE FROM attendance WHERE student_id = ? AND attendance_date = ? AND attendance_time = ?", (student_id, date, time))
            
            self.load_data()
            messagebox.showinfo("Thành công", f"Đã hủy thành công {len(valid_records)} lượt điểm danh.")
        except Exception as e:
            messagebox.showerror("Lỗi", f"Không thể xóa: {e}")

    def select_all(self, event=None):
        """Hàm tự động bôi đen toàn bộ các dòng đang hiển thị trên bảng"""
        all_items = self.tree.get_children()
        
        if all_items:
            self.tree.selection_set(all_items)
            
        return "break"

    # ==========================================
    # CÁC HÀM XUẤT FILE 
    # ==========================================
    def export_excel(self):
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel Files", "*.xlsx")], initialfile="lich_su_diem_danh.xlsx")
        if not path: return
        
        try:
            import pandas as pd
            from openpyxl.styles import Font, Alignment, PatternFill
            
            data = [self.tree.item(child)["values"] for child in self.tree.get_children()]
            df = pd.DataFrame(data, columns=["MSSV", "Họ tên", "Lớp", "Môn học", "Ngày", "Giờ", "Trạng thái"])
            
            with pd.ExcelWriter(path, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name='Lich_Su')
                worksheet = writer.sheets['Lich_Su']
                
                widths = {'A': 12, 'B': 25, 'C': 10, 'D': 25, 'E': 12, 'F': 12, 'G': 15}
                for col, width in widths.items():
                    worksheet.column_dimensions[col].width = width
                    
                header_fill = PatternFill(start_color="107c41", end_color="107c41", fill_type="solid")
                for cell in worksheet["1:1"]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.fill = header_fill
                    
                for row in worksheet.iter_rows(min_row=2):
                    for idx, cell in enumerate(row):
                        if idx not in [1, 3]: 
                            cell.alignment = Alignment(horizontal="center")
                            
            messagebox.showinfo("Thành công", "Đã xuất file Excel phiên bản chuyên nghiệp thành công!")
        except ImportError:
            messagebox.showerror("Thiếu thư viện", "Vui lòng cài đặt: pip install pandas openpyxl")
        except Exception as e:
            messagebox.showerror("Lỗi", f"Không thể xuất file: {e}")

    def export_pdf(self):
        path = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF Files", "*.pdf")], initialfile="lich_su_diem_danh.pdf")
        if not path: return
        
        try:
            from fpdf import FPDF
            import os
            
            pdf = FPDF()
            pdf.add_page()
            
            font_path = "C:\\Windows\\Fonts\\arial.ttf"
            has_unicode = os.path.exists(font_path)
            
            if has_unicode:
                pdf.add_font("ArialVN", "", font_path)
                pdf.set_font("ArialVN", "", 18)
            else:
                pdf.set_font("Helvetica", "B", 18)
                
            pdf.cell(0, 15, "BÁO CÁO ĐIỂM DANH SINH VIÊN", new_x="LMARGIN", new_y="NEXT", align="C")
            
            if has_unicode:
                pdf.set_font("ArialVN", "", 10)
            else:
                pdf.set_font("Helvetica", "", 10)
            pdf.cell(0, 10, f"Thời điểm xuất file: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", new_x="LMARGIN", new_y="NEXT", align="R")
            pdf.ln(5)
            
            if has_unicode:
                pdf.set_font("ArialVN", "", 11)
            
            headers = ["MSSV", "Họ tên", "Lớp", "Môn học", "Ngày", "Giờ", "Trạng thái"]
            col_widths = [20, 45, 15, 45, 25, 20, 25]
            
            for i, header in enumerate(headers):
                pdf.cell(col_widths[i], 10, header, border=1, align="C")
            pdf.ln()
            
            for child in self.tree.get_children():
                values = self.tree.item(child)["values"]
                for i, val in enumerate(values):
                    align = "L" if i in [1, 3] else "C" 
                    pdf.cell(col_widths[i], 10, str(val), border=1, align=align)
                pdf.ln()
                
            pdf.output(path)
            messagebox.showinfo("Thành công", "Đã xuất báo cáo ra file PDF thành công!")
            
        except ImportError:
            messagebox.showerror("Thiếu thư viện", "Vui lòng cài đặt: pip install fpdf2")
        except Exception as e:
            messagebox.showerror("Lỗi", f"Không thể xuất file: {e}")